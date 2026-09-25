"""
Builds the complete Summer Training & MedReport AI presentation PPTX
for Anirudh Varshney (JUIT, Roll No: 241030149).
Modifies the official JUIT presentation template preserving all logos,
layouts, headers, styling, and visual theme.
"""

import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Fix utf-8 output
sys.stdout.reconfigure(encoding="utf-8")

file_path = r"C:\Users\421je\.gemini\antigravity\scratch\medical-report-summarizer\Anirudh_Varshney_Summer_Training_Presentation.pptx"
prs = pptx.Presentation(file_path)

def set_text_safely(shape, new_text, font_size=None, bold=None, color=None):
    """Sets text in a shape while preserving paragraph and font formatting if specified."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    
    # Store existing font properties from first run if available
    old_font_name = None
    old_font_size = None
    old_font_bold = None
    old_font_color = None
    
    if tf.paragraphs and tf.paragraphs[0].runs:
        run0 = tf.paragraphs[0].runs[0]
        old_font_name = run0.font.name
        old_font_size = run0.font.size
        old_font_bold = run0.font.bold
        try:
            if run0.font.color and run0.font.color.rgb:
                old_font_color = run0.font.color.rgb
        except Exception:
            pass

    tf.text = ""  # Clear
    lines = new_text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        
        # Apply font name
        if old_font_name:
            run.font.name = old_font_name
            
        # Apply font size
        if font_size:
            run.font.size = Pt(font_size)
        elif old_font_size:
            run.font.size = old_font_size
            
        # Apply bold
        if bold is not None:
            run.font.bold = bold
        elif old_font_bold is not None:
            run.font.bold = old_font_bold
            
        # Apply color
        if color:
            run.font.color.rgb = color
        elif old_font_color:
            run.font.color.rgb = old_font_color

# =========================================================================
# SLIDE 1: Title Slide
# =========================================================================
s1 = prs.slides[0]
for shape in s1.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Ipshit Singh Rana" in t:
            set_text_safely(shape, "Anirudh Varshney")
        elif "241030331" in t or "241030149" in t:
            set_text_safely(shape, "Roll No.: 241030149\nB.Tech CSE, Third Year, JUIT")
        elif "IBM/Coursera AI Foundations" in t:
            set_text_safely(shape, "Microsoft Certified: Azure AI Fundamentals (AI-901)")
        elif "Course Evaluation" in t:
            set_text_safely(shape, "Summer Training & Capstone Project Evaluation")
        elif "4 Courses" in t or "6 Weeks" in t:
            set_text_safely(shape, "Microsoft Learn | 6 Weeks Summer Training\nCredential ID: CB867A28B596E933 | Earned: 15 July 2026")

# =========================================================================
# SLIDE 2: What I Learned (Roadmap)
# =========================================================================
s2 = prs.slides[1]
for shape in s2.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "AI Fundamentals" in t and "Understand what AI is" in t:
            set_text_safely(shape, "Azure AI Fundamentals\nUnderstand cloud AI & services")
        elif "Generative AI" in t and "Create new content" in t:
            set_text_safely(shape, "Computer Vision & OCR\nAzure Document Intelligence")
        elif "Prompt Engineering" in t and "Ask AI better questions" in t:
            set_text_safely(shape, "Generative AI & LLMs\nPrompting, RAG & Foundry")
        elif "AI Chatbots" in t and "Use actions and context" in t:
            set_text_safely(shape, "Applied Capstone\nMedReport AI on Azure")
        elif "The courses moved from" in t:
            set_text_safely(shape, "The 6-week training progressed from Azure cloud fundamentals to Computer Vision, Generative AI, and building MedReport AI.")

# =========================================================================
# SLIDE 3: Introduction to Artificial Intelligence & Azure AI
# =========================================================================
s3 = prs.slides[2]
for shape in s3.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Introduction to Artificial Intelligence" in t:
            set_text_safely(shape, "Introduction to Artificial Intelligence & Azure AI")
        elif "Artificial Intelligence" in t and "Computers performing tasks" in t:
            set_text_safely(shape, "Artificial Intelligence\nComputers performing cognitive tasks that normally require human intelligence.")
        elif "Augmented Intelligence" in t:
            set_text_safely(shape, "Azure AI Services\nPre-built, scalable cloud APIs for Vision, Speech, Language, and Decision.")
        elif "Real-life examples" in t:
            set_text_safely(shape, "Real-life Azure AI Applications")
        elif "Recommendation systems suggest" in t:
            set_text_safely(shape, "•  Azure Document Intelligence extracts structured data from medical files.\n•  Azure OpenAI powers conversational summarization and clinical reasoning.\n•  Azure Machine Learning trains and deploys custom diagnostic models.")

# =========================================================================
# SLIDE 4: AI, ML and Deep Learning on Azure
# =========================================================================
s4 = prs.slides[3]
for shape in s4.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "AI, ML and Deep Learning" in t:
            set_text_safely(shape, "AI, ML and Deep Learning on Azure")
        elif t == "Machine Learning\nLearns from data" or ("Machine Learning" in t and "Learns from data" in t):
            set_text_safely(shape, "Machine Learning\nLearns patterns from data (Azure ML)")
        elif t == "Deep Learning\nLearns complex patterns" or ("Deep Learning" in t and "Learns complex patterns" in t):
            set_text_safely(shape, "Deep Learning\nNeural networks for vision & language")
        elif "Foundation Model" in t:
            set_text_safely(shape, "Foundation Models\nPre-trained models supporting diverse cloud workloads.")
        elif "Large Language Model" in t:
            set_text_safely(shape, "Large Language Models (LLMs)\nAdvanced models like GPT-4o deployed via Azure OpenAI.")
        elif "Simple relationship:" in t:
            set_text_safely(shape, "Relationship: Deep Learning powers modern LLMs and Cognitive Services, which are part of AI on Microsoft Azure.")

# =========================================================================
# SLIDE 5: How Azure AI Understands the World
# =========================================================================
s5 = prs.slides[4]
for shape in s5.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "How AI Understands the World" in t:
            set_text_safely(shape, "How Azure AI Understands the World")
        elif "NLP" in t and "Understands text and language" in t:
            set_text_safely(shape, "Natural Language Processing (NLP)\nAnalyzes and extracts meaning from medical text.")
        elif "Example: summarizing a message" in t:
            set_text_safely(shape, "Example: Clinical report summarization")
        elif "Speech Recognition" in t:
            set_text_safely(shape, "Speech & Audio AI\nConverts spoken words to text and generates natural audio.")
        elif "Example: voice typing" in t:
            set_text_safely(shape, "Example: Voice-assisted doctor briefings")
        elif "Computer Vision" in t:
            set_text_safely(shape, "Computer Vision & OCR\nInterprets visual documents, PDFs, and diagnostic scans.")
        elif "Example: detecting an object in a photo" in t:
            set_text_safely(shape, "Example: Azure Document Intelligence")

# =========================================================================
# SLIDE 6: Responsible AI & Governance on Azure
# =========================================================================
s6 = prs.slides[5]
for shape in s6.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "AI in the Real World" in t:
            set_text_safely(shape, "Responsible AI & Governance on Azure")
        elif "IoT" in t and "Connected devices" in t:
            set_text_safely(shape, "Fairness & Inclusivity\nAccessible to patients of all backgrounds.")
        elif "Edge Computing" in t:
            set_text_safely(shape, "Reliability & Safety\nZero-hallucination validation for medical data.")
        elif "AI Ethics" in t:
            set_text_safely(shape, "Privacy & Security\nProtected health data via Azure compliance.")
        elif "Governance" in t and "Rules for responsible" in t:
            set_text_safely(shape, "Transparency & Audit\nAuditable raw data alongside AI conclusions.")
        elif "Example: Smart camera" in t:
            set_text_safely(shape, "Responsible AI in Action: Ingests clinical report → sanitizes data → validates lab boundaries → presents transparent summary.")
        elif "•  IoT connects" in t:
            set_text_safely(shape, "•  Microsoft Responsible AI standard guides ethical design across all 6 core pillars.\n•  Clinical guardrails and reference grounding eliminate harmful medical hallucinations.\n•  Auditable data retention ensures patient privacy and regulatory compliance.")

# =========================================================================
# SLIDE 7: Traditional AI vs Generative AI in Healthcare (Table)
# =========================================================================
s7 = prs.slides[6]
for shape in s7.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Traditional AI vs Generative AI" in t:
            set_text_safely(shape, "Traditional AI vs Generative AI in Healthcare")
    elif shape.has_table:
        table = shape.table
        # Update table rows
        # Row 0: Header
        table.cell(0, 0).text = "Aspect"
        table.cell(0, 1).text = "Traditional / Deterministic AI"
        table.cell(0, 2).text = "Generative AI"
        
        # Row 1: Purpose
        table.cell(1, 0).text = "Main purpose"
        table.cell(1, 1).text = "Evaluates boundaries & flags anomalies"
        table.cell(1, 2).text = "Synthesizes plain-language summaries"
        
        # Row 2: Output
        table.cell(2, 0).text = "Typical output"
        table.cell(2, 1).text = "Categorical label: High, Low, Critical"
        table.cell(2, 2).text = "Narrative explanations & recommendations"
        
        # Row 3: Example
        table.cell(3, 0).text = "Healthcare example"
        table.cell(3, 1).text = "Flagging HbA1c > 5.7% (Prediabetes)"
        table.cell(3, 2).text = "Explaining dietary & lifestyle adjustments"
        
        # Row 4: Strength
        table.cell(4, 0).text = "Core strength"
        table.cell(4, 1).text = "100% precision & zero hallucination"
        table.cell(4, 2).text = "Fluent, patient-friendly communication"

# =========================================================================
# SLIDE 8: Generative AI and Advanced AI Types
# =========================================================================
s8 = prs.slides[7]
for shape in s8.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Generative AI and Advanced AI Types" in t:
            set_text_safely(shape, "Generative AI and Advanced AI Types")
        elif "Creates new content from a user request" in t:
            set_text_safely(shape, "Generative AI\nSynthesizes natural clinical explanations from complex diagnostic data.")
        elif "Multimodal AI" in t and "Works with more than one" in t:
            set_text_safely(shape, "Multimodal AI\nProcesses diagnostic images, tabular charts, and narrative text simultaneously.")
        elif "Agentic AI" in t and "Can perform multiple steps" in t:
            set_text_safely(shape, "Agentic AI\nOrchestrates multi-step pipelines: upload → extract → validate → summarize → store.")
        elif "Easy examples" in t:
            set_text_safely(shape, "Practical Healthcare Implementations")
        elif "•  Generative: create a short email" in t:
            set_text_safely(shape, "•  Generative: Translates complex blood panels into 3-5 sentence patient summaries.\n•  Multimodal: Azure Document Intelligence reading layout + text + embedded tables.\n•  Agentic: End-to-end autonomous report analysis in MedReport AI.")

# =========================================================================
# SLIDE 9: Prompt Engineering & Clinical Grounding
# =========================================================================
s9 = prs.slides[8]
for shape in s9.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Prompt Engineering" in t and shape.name.startswith("Google"):
            set_text_safely(shape, "Prompt Engineering & Clinical Grounding")
        elif "Prompt\nThe instruction or question" in t or (t.startswith("Prompt") and "instruction" in t):
            set_text_safely(shape, "Prompt\nInput instruction provided to guide generative model reasoning.")
        elif "Prompt Engineering\nWriting prompts clearly" in t or ("Writing prompts clearly" in t):
            set_text_safely(shape, "Prompt Engineering\nDesigning structured prompts to ensure clinical accuracy and JSON output.")
        elif "Why it matters\nClear input" in t or ("Why it matters" in t):
            set_text_safely(shape, "Why It Matters\nPrecise clinical constraints prevent medical hallucinations.")
        elif "Bad prompt" in t or "Tell me about AI" in t:
            set_text_safely(shape, "Ambiguous Prompt\n“Tell me what is wrong with this medical report.”")
        elif "Better prompt" in t or "Explain AI to a first-year" in t:
            set_text_safely(shape, "Engineered Clinical Prompt\n“Act as a clinical pathologist. Extract all lab parameters, match against clinical reference ranges, and return structured JSON with severity and recommendations.”")

# =========================================================================
# SLIDE 10: Prompting Techniques & Microsoft Foundry
# =========================================================================
s10 = prs.slides[9]
for shape in s10.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Prompting Techniques" in t:
            set_text_safely(shape, "Prompting Techniques & Microsoft Foundry")
        elif "Zero-shot" in t:
            set_text_safely(shape, "Zero-shot\nDirect query without previous examples.")
        elif "Few-shot" in t:
            set_text_safely(shape, "Few-shot\nProviding exemplar lab formats.")
        elif "Persona" in t:
            set_text_safely(shape, "Persona Role\nInstructing AI to act as medical analyst.")
        elif "Interview" in t:
            set_text_safely(shape, "RAG Grounding\nGrounding responses in verified medical references.")
        elif "Chain-of-Thought" in t:
            set_text_safely(shape, "Chain-of-Thought\nStep-by-step reasoning for lab panels.")
        elif "Tree-of-Thought" in t:
            set_text_safely(shape, "System Guardrails\nEnforcing valid JSON schemas.")

# =========================================================================
# SLIDE 11: Applied Capstone: MedReport AI
# =========================================================================
s11 = prs.slides[10]
for shape in s11.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "AI-Powered Chatbots" in t or "Title 1" in shape.name:
            set_text_safely(shape, "Applied Capstone: MedReport AI")
        elif "Decision-Tree Chatbot" in t:
            set_text_safely(shape, "The Real-World Problem\nComplex lab jargon, dense tables, patient anxiety, and missed early medical alerts.")
        elif "Generative AI Chatbot" in t:
            set_text_safely(shape, "The Cloud AI Solution\nMedReport AI: Cloud-native medical report analyzer powered by Microsoft Azure.")
        elif "Customer-support example" in t:
            set_text_safely(shape, "Core System Capabilities")
        elif "•  Intent → what the user wants" in t:
            set_text_safely(shape, "•  Instant Ingestion: Supports multi-page PDFs, clinical scans, and lab images.\n•  Automatic Anomaly Detection: Detects high/low/critical lab values across 16+ panels.\n•  Interactive Dashboard: Color-coded severity badges with doctor-ready recommendations.")

# =========================================================================
# SLIDE 12: Azure Cloud Services Architecture
# =========================================================================
s12 = prs.slides[11]
for shape in s12.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "watsonx Assistant" in t or "Title 1" in shape.name:
            set_text_safely(shape, "Azure Cloud Services Architecture")
        elif "User" in t and "I want to track" in t:
            set_text_safely(shape, "User Upload\nMedical PDF / Diagnostic Scan")
        elif "Intent" in t and "Track Order" in t:
            set_text_safely(shape, "Azure Storage\nBlob Storage (`anirudhmedreport`)")
        elif "Required info" in t and "Order Number" in t:
            set_text_safely(shape, "Azure Vision\nDocument Intelligence (`prebuilt-layout`)")
        elif "Workflow building blocks" in t:
            set_text_safely(shape, "Architectural Resilience & Failover")
        elif "•  Actions → steps the chatbot" in t:
            set_text_safely(shape, "•  Azure Document Intelligence: Extracts multi-column tables, text, and key-value pairs.\n•  Azure Blob Storage: Secure cloud persistence with unique UUID isolation.\n•  Dual-Engine Fallback: Auto-routes files >4MB to high-speed stream parser (bypassing F0 limits).\n•  Deterministic Rule Engine: Guarantees 100% clinical accuracy with zero API quota dependency.")

# =========================================================================
# SLIDE 13: Complete AI Workflow & Validation
# =========================================================================
s13 = prs.slides[12]
for shape in s13.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Complete AI Workflow" in t or "Title 1" in shape.name:
            set_text_safely(shape, "Complete MedReport AI Workflow")
        elif t == "USER":
            set_text_safely(shape, "USER")
        elif t == "INPUT":
            set_text_safely(shape, "PDF FILE")
        elif t == "AI / ML":
            set_text_safely(shape, "AZURE BLOB")
        elif t == "GENERATIVE AI":
            set_text_safely(shape, "AZURE OCR")
        elif t == "PROMPT":
            set_text_safely(shape, "CLINICAL RULE")
        elif t == "CHATBOT ACTION":
            set_text_safely(shape, "JSON SCHEMA")
        elif t == "RESPONSE":
            set_text_safely(shape, "DASHBOARD")
        elif "Example: customer asks about an order" in t:
            set_text_safely(shape, "Live Validation on 8.3 MB, 12-page Report: Successfully identified Critical TSH (17.0 uIU/mL), High HbA1c (6.2%), High Triglycerides (199 mg/dL), and generated 8 targeted clinical recommendations.")
        elif "Connected learning" in t or "AI fundamentals →" in t:
            set_text_safely(shape, "Connected Pipeline: User Upload → Azure Cloud Ingestion → Structural OCR → Clinical Evaluation → Interactive Visual Dashboard")

# =========================================================================
# SLIDE 14: Key Learnings
# =========================================================================
s14 = prs.slides[13]
for shape in s14.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Key Learnings" in t and shape.name.startswith("Title"):
            set_text_safely(shape, "Key Learnings & Training Outcomes")
        elif "•  AI enables computers to perform" in t:
            set_text_safely(shape, "•  Mastered core AI workloads: Computer Vision, Natural Language Processing, and Generative AI.\n•  Hands-on cloud engineering with Microsoft Azure AI Services, Microsoft Foundry, and Azure ML Studio.\n•  Understood and implemented Retrieval-Augmented Generation (RAG) to ground AI in verified data.\n•  Architected an end-to-end production web application (MedReport AI) solving healthcare challenges.\n•  Engineered resilient failover systems to handle cloud tier limits (4MB threshold) and API quotas.\n•  Applied Microsoft's Responsible AI framework: prioritizing patient safety, privacy, and transparency.\n•  Successfully earned the Microsoft Certified: Azure AI Fundamentals credential.")

# =========================================================================
# SLIDE 15: Conclusion & Career Roadmap
# =========================================================================
s15 = prs.slides[14]
for shape in s15.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text
        if "Conclusion" in t and shape.name.startswith("Title"):
            set_text_safely(shape, "Conclusion & Future Roadmap")
        elif "What I gained from the four courses" in t or "I learned the fundamentals" in t:
            set_text_safely(shape, "What I Gained from the 6-Week Summer Training\nSuccessfully bridged foundational academic theory with enterprise cloud development on Microsoft Azure, delivering a fully functional applied AI capstone project.")
        elif "•  Understand the fundamentals" in t:
            set_text_safely(shape, "•  Master Azure AI Workloads  •  Deploy Cloud AI Services  •  Build Resilient Solutions")
        elif "Thank You" in t:
            set_text_safely(shape, "Thank You!\nAnirudh Varshney (Roll No: 241030149, JUIT)")

# Save final presentation
prs.save(file_path)
print("Complete presentation generated successfully at:", file_path)
